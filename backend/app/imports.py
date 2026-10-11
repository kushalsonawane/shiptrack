import io
import re
from datetime import date, datetime

import openpyxl

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
GSTIN_RE = re.compile(r"^\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]$")
PAN_RE = re.compile(r"^[A-Z]{5}\d{4}[A-Z]$")
PIN_RE = re.compile(r"^\d{6}$")
PHONE_RE = re.compile(r"^[+]?[\d][\d\s\-()/]{7,19}$")

INDIAN_STATES = {
    "andhra pradesh", "arunachal pradesh", "assam", "bihar", "chhattisgarh", "goa",
    "gujarat", "haryana", "himachal pradesh", "jharkhand", "karnataka", "kerala",
    "madhya pradesh", "maharashtra", "manipur", "meghalaya", "mizoram", "nagaland",
    "odisha", "punjab", "rajasthan", "sikkim", "tamil nadu", "telangana", "tripura",
    "uttar pradesh", "uttarakhand", "west bengal", "delhi", "chandigarh",
    "jammu and kashmir", "ladakh", "puducherry", "dadra and nagar haveli",
    "daman and diu", "andaman and nicobar islands", "lakshadweep",
}

STATE_ALIASES = {
    "gujrat": "Gujarat",
    "maharastra": "Maharashtra",
    "maharashra": "Maharashtra",
    "orissa": "Odisha",
    "uttaranchal": "Uttarakhand",
    "up": "Uttar Pradesh",
    "mp": "Madhya Pradesh",
    "ap": "Andhra Pradesh",
    "benglore": "Karnataka",
    "karnatka": "Karnataka",
    "hariyana": "Haryana",
    "d & nh": "Dadra and Nagar Haveli",
    "daman": "Daman and Diu",
}

HEADER_ALIASES = {
    "legal_name": ("company name", "legal name", "party name"),
    "gstin": ("gst no", "gstin", "gst number", "gst"),
    "pan": ("pan no", "pan number", "pan"),
    "email": ("email", "email id", "email address"),
    "phone": ("person phone number", "phone number", "phone", "mobile", "contact number"),
    "state": ("shipping address state", "billing address state", "state"),
    "city": ("shipping address city", "billing address city", "city"),
    "pincode": ("shipping address postal code", "billing address postal code", "pincode", "pin code", "postal code"),
    "address1": ("shipping address line 1", "billing address line 1", "address line 1", "address1"),
    "address2": ("shipping address line 2", "billing address line 2", "address line 2", "address2"),
    "payment_term": ("payment term", "payment terms"),
    "tags": ("tags",),
    "gst_type": ("gst type",),
    "country": ("shipping address country", "billing address country", "country"),
    "contact_first": ("customer contact detail 1 - first name", "vendor contact detail 1 - first name"),
    "contact_last": ("customer contact detail 1 - last name", "vendor contact detail 1 - last name"),
    "contact_email": ("customer contact detail 1 - email", "vendor contact detail 1 - email"),
    "contact_phone": ("customer contact detail 1 - phone number", "vendor contact detail 1 - phone number"),
}


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _looks_like_name(text: str) -> bool:
    upper = text.upper()
    if EMAIL_RE.match(text) or GSTIN_RE.match(upper) or PAN_RE.match(upper):
        return False
    if re.match(r"^[\d\s\-/+,.;]+$", text):
        return False
    if text.lower() in INDIAN_STATES or text.lower() in STATE_ALIASES:
        return False
    return len(text) >= 3 and any(ch.isalpha() for ch in text)


def _match_state(text: str):
    lowered = text.strip().lower()
    if lowered in INDIAN_STATES:
        return text.strip().title()
    if lowered in STATE_ALIASES:
        return STATE_ALIASES[lowered]
    return None


def _build_colmap(header_row) -> dict:
    colmap = {}
    for index, cell in enumerate(header_row):
        key = _clean(cell).lower()
        if not key:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field not in colmap and key in aliases:
                colmap[field] = index
    return colmap


def _apply_heuristics(record: dict, row) -> dict:
    for cell in row:
        text = _clean(cell)
        if not text:
            continue
        upper = text.upper()
        if not record.get("gstin") and GSTIN_RE.match(upper):
            record["gstin"] = upper
            continue
        if not record.get("pan") and PAN_RE.match(upper):
            record["pan"] = upper
            continue
        if not record.get("email") and EMAIL_RE.match(text):
            record["email"] = text.lower()
            continue
        if not record.get("state"):
            state = _match_state(text)
            if state:
                record["state"] = state
                continue
        if not record.get("pincode") and PIN_RE.match(text):
            record["pincode"] = text
            continue
        if not record.get("phone") and PHONE_RE.match(re.sub(r"\s?[Ee]xt\.?\s?\d+$", "", text)):
            digits = re.sub(r"\D", "", text)
            if 8 <= len(digits) <= 13:
                record["phone"] = text
                continue
        if not record.get("legal_name") and _looks_like_name(text):
            record["legal_name"] = text
            continue
        if not record.get("address1") and len(text) > 25 and any(ch.isalpha() for ch in text):
            record["address1"] = text
    return record


def _row_to_record(row, colmap: dict) -> dict:
    record = {}
    for field in (
        "legal_name", "gstin", "pan", "email", "phone", "state", "city", "pincode",
        "address1", "address2", "payment_term", "tags", "gst_type", "country",
        "contact_first", "contact_last", "contact_email", "contact_phone",
    ):
        index = colmap.get(field)
        record[field] = _clean(row[index]) if index is not None and index < len(row) else ""

    if record["gstin"] and not GSTIN_RE.match(record["gstin"].upper()):
        record["gstin"] = ""
    if record["pan"] and not PAN_RE.match(record["pan"].upper()):
        record["pan"] = ""
    if record["email"] and not EMAIL_RE.match(record["email"]):
        record["email"] = ""
    if record["pincode"] and not PIN_RE.match(record["pincode"]):
        record["pincode"] = ""
    if record["phone"]:
        digits = re.sub(r"\D", "", record["phone"])
        if not (8 <= len(digits) <= 15):
            record["phone"] = ""
    if record["contact_email"] and not EMAIL_RE.match(record["contact_email"]):
        record["contact_email"] = ""

    return _apply_heuristics(record, row)


def _parse_sheet(rows):
    if not rows:
        return [], []
    header_index = None
    for index, row in enumerate(rows[:10]):
        joined = " ".join(_clean(cell).lower() for cell in row)
        if "company name" in joined or ("gst" in joined and "email" in joined):
            header_index = index
            break

    colmap = _build_colmap(rows[header_index]) if header_index is not None else {}
    data_rows = rows[header_index + 1:] if header_index is not None else rows

    records, errors = [], []
    for offset, row in enumerate(data_rows):
        row_number = (header_index or 0) + 2 + offset
        if not any(_clean(cell) for cell in row):
            continue
        record = _row_to_record(row, colmap)
        if len(record.get("legal_name", "")) >= 2:
            records.append({key: (value or None) for key, value in record.items()})
        else:
            errors.append({"row": row_number, "reason": "no company name found"})
    return records, errors


def parse_contacts_xlsx(data: bytes):
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    records, errors = [], []
    for sheet_name in workbook.sheetnames:
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        sheet_records, sheet_errors = _parse_sheet(rows)
        for error in sheet_errors:
            error["sheet"] = sheet_name
        records.extend(sheet_records)
        errors.extend(sheet_errors)
    if not records and not errors:
        errors.append({"sheet": "-", "row": 0, "reason": "empty workbook"})
    return records, errors


INVOICE_HEADER_ALIASES = {
    "invoice_no": ("invoice no", "invoice no.", "invoice number"),
    "invoice_date": ("invoice date", "date"),
    "direction": ("direction",),
    "party_gstin": ("party gstin", "gst no", "gstin", "gst number"),
    "party_name": ("party name", "customer", "vendor", "party"),
    "po_no": ("po no", "po no.", "po number", "po"),
    "description": ("description", "item", "item description"),
    "hsn_sac": ("hsn", "hsn/sac", "hsn sac"),
    "quantity": ("qty", "quantity"),
    "unit": ("unit", "uom"),
    "rate": ("rate", "unit price", "price"),
    "gst_rate": ("gst rate", "gst %", "gst%", "tax %"),
}


def _to_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _to_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return default


def parse_invoices_xlsx(data: bytes):
    """Parse an invoice import sheet.

    One row per line item; rows are grouped by invoice number. Required columns:
    Invoice No; Party is resolved from Party GSTIN or Party Name.
    """
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        return [], [{"row": 0, "reason": "empty sheet"}]

    header_index = None
    for index, row in enumerate(rows[:10]):
        joined = " ".join(_clean(cell).lower() for cell in row)
        if "invoice" in joined and ("party" in joined or "customer" in joined or "vendor" in joined):
            header_index = index
            break
    if header_index is None:
        return [], [{"row": 0, "reason": "header row with Invoice No and Party not found"}]

    colmap = {}
    for index, cell in enumerate(rows[header_index]):
        key = _clean(cell).lower()
        for field, aliases in INVOICE_HEADER_ALIASES.items():
            if field not in colmap and key in aliases:
                colmap[field] = index
    if "invoice_no" not in colmap:
        return [], [{"row": header_index + 1, "reason": "Invoice No column not found"}]

    grouped = {}
    errors = []
    for offset, row in enumerate(rows[header_index + 1:]):
        row_number = header_index + 2 + offset
        if not any(_clean(cell) for cell in row):
            continue

        def cell(field):
            index = colmap.get(field)
            return row[index] if index is not None and index < len(row) else None

        invoice_no = _clean(cell("invoice_no"))
        if not invoice_no:
            errors.append({"row": row_number, "reason": "missing invoice number"})
            continue
        entry = grouped.setdefault(invoice_no, {"rows": []})
        entry["rows"].append(
            {
                "party_gstin": _clean(cell("party_gstin")).upper() or None,
                "party_name": _clean(cell("party_name")) or None,
                "invoice_date": _to_date(cell("invoice_date")),
                "direction": (_clean(cell("direction")) or "").lower() or None,
                "po_no": _clean(cell("po_no")) or None,
                "description": _clean(cell("description")) or None,
                "hsn_sac": _clean(cell("hsn_sac")) or None,
                "quantity": _to_float(cell("quantity"), 1.0),
                "unit": _clean(cell("unit")) or "Nos",
                "rate": _to_float(cell("rate"), 0.0),
                "gst_rate": _to_float(cell("gst_rate"), 18.0),
            }
        )

    records = []
    for invoice_no, entry in grouped.items():
        first = entry["rows"][0]
        records.append(
            {
                "invoice_no": invoice_no,
                "party_gstin": first["party_gstin"],
                "party_name": first["party_name"],
                "invoice_date": first["invoice_date"],
                "direction": first["direction"],
                "po_no": first["po_no"],
                "items": [
                    {
                        "description": row["description"],
                        "hsn_sac": row["hsn_sac"],
                        "quantity": row["quantity"],
                        "unit": row["unit"],
                        "rate": row["rate"],
                        "gst_rate": row["gst_rate"],
                    }
                    for row in entry["rows"]
                    if row["description"]
                ],
            }
        )
    return records, errors
